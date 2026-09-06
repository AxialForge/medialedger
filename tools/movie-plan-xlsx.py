"""Build an Excel report: current movie file names vs what the naming engine proposes.
Usage: python tools/movie-plan-xlsx.py <plan.json> <out.xlsx>
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo

plan = json.load(open(sys.argv[1], encoding='utf8'))
out = sys.argv[2]

F = 'Arial'
font = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color='FFFFFF')
hdr_fill = PatternFill('solid', fgColor='1F3864')
title_font = Font(name=F, size=14, bold=True)
grey = Font(name=F, size=9, color='666666')
thin = Side(style='thin', color='BFBFBF')

FLAG_TEXT = {
    'no_source': 'No LiLTV / WEB / rip marker in the old name -> "Source" placeholder. Fill in via Fix... (Source field).',
    'no_year': 'No release year found -> "(Year)" placeholder. Plex will not match the file until it is set via Fix....',
    'res_mismatch': 'Old name claimed one resolution, ffprobe measured another. The probed value is used.',
    'hdr_uncertain': 'BT.2020 colour primaries without an HDR transfer curve; treated as HDR.',
    'hdr_claimed_but_sdr': 'Old name says HDR / Dolby Vision but the probe found SDR. SDR is used.',
    'audio_und': 'Audio track has no language tag ("und"); audio token omitted.',
    'audio_partly_und': 'Some audio tracks have no language tag.',
    'audio_unknown': 'No audio language data at all.',
    'collision': 'Two or more files would receive the identical name; all of them are blocked until one is fixed or removed.',
}

def status_of(p):
    if p.get('unchanged'): return 'unchanged'
    return 'ready' if p.get('ok') else 'blocked'

wb = Workbook()

# ---------------- Plan sheet ----------------
ws = wb.active
ws.title = 'Plan'
headers = ['Status', 'Current name', 'Proposed name', 'Flags', 'Blocked reason', 'Title', 'Year', 'Source', 'Resolution', 'HDR', 'Codec', 'Audio', 'Edition',
           'Probed size (GB)', 'Probed WxH', 'Probed HDR', 'Probed codec', 'Probed audio', 'Bitrate (kbps)', 'Length (min)', 'Folder']
ws.append(headers)
for p in plan:
    t = p.get('tokens') or {}
    flags = ' '.join(p.get('flags') or [])
    if p.get('blocked') and 'collision' in p['blocked']: flags = (flags + ' collision').strip()
    ws.append([
        status_of(p), p['from'], p.get('name') or '', flags, p.get('blocked') or '',
        t.get('title') or '', t.get('year') or '', t.get('source') or '', t.get('resolution') or '', t.get('hdr') or '', t.get('codec') or '', t.get('audio') or '', t.get('edition') or '',
        round((p.get('size') or 0) / 1073741824, 2), f"{p.get('width') or ''}x{p.get('height') or ''}" if p.get('width') else '', p.get('hdrRaw') or '', p.get('codecRaw') or '', p.get('audioRaw') or '',
        p.get('bitrate') or '', round((p.get('duration_s') or 0) / 60, 1) if p.get('duration_s') else '', p.get('dir') or '',
    ])
n = ws.max_row
for c in ws[1]:
    c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2, max_row=n):
    for c in row: c.font = font; c.alignment = Alignment(vertical='top')
widths = [10, 58, 62, 26, 34, 34, 7, 8, 10, 6, 8, 10, 30, 13, 11, 12, 12, 12, 12, 11, 30]
for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{n}'
ws.row_dimensions[1].height = 30
ready_fill = PatternFill('solid', fgColor='E2F0D9'); blocked_fill = PatternFill('solid', fgColor='F8CBAD'); flag_fill = PatternFill('solid', fgColor='FFF2CC'); same_fill = PatternFill('solid', fgColor='EDEDED')
rng = f'A2:{get_column_letter(len(headers))}{n}'
ws.conditional_formatting.add(rng, FormulaRule(formula=['$A2="blocked"'], fill=blocked_fill))
ws.conditional_formatting.add(rng, FormulaRule(formula=['$A2="unchanged"'], fill=same_fill))
ws.conditional_formatting.add(f'D2:D{n}', FormulaRule(formula=['LEN($D2)>0'], fill=flag_fill))
ws.conditional_formatting.add(f'A2:A{n}', FormulaRule(formula=['$A2="ready"'], fill=ready_fill))
for col in 'GHIJKL':
    ws.conditional_formatting.add(f'{col}2:{col}{n}', FormulaRule(formula=[f'OR(${col}2="Year",${col}2="Source")'], fill=flag_fill, font=Font(name=F, size=10, bold=True, color='9C5700')))

# ---------------- Summary sheet (formulas over Plan) ----------------
s = wb.create_sheet('Summary', 0)
s['A1'] = 'MediaLedger - Movie rename plan'; s['A1'].font = title_font
s['A2'] = 'Current file names vs. the names the naming engine would apply. Generated from the MediaLedger database; nothing on the share has been changed.'; s['A2'].font = grey
s['A3'] = 'Pattern: Title (Year) - Source Resolution HDR Codec [Audio] [{edition-Name}].ext   -   Resolution, HDR, codec and audio come only from ffprobe; Year/Source placeholders mark values the engine could not prove.'; s['A3'].font = grey
r = 5
def kv(label, formula, note=''):
    global r
    s.cell(r, 1, label).font = bold; c = s.cell(r, 2, formula); c.font = font; c.alignment = Alignment(horizontal='right')
    if note: s.cell(r, 3, note).font = grey
    r += 1
kv('Movie files in plan', f'=COUNTA(Plan!A2:A{n})')
kv('Ready to rename', f'=COUNTIF(Plan!A2:A{n},"ready")', 'proposed name differs and every safety check passes')
kv('Already correct', f'=COUNTIF(Plan!A2:A{n},"unchanged")')
kv('Blocked', f'=COUNTIF(Plan!A2:A{n},"blocked")', 'never renamed until fixed - see Blocked sheet')
kv('   of which name collisions', f'=COUNTIF(Plan!E2:E{n},"collision*")')
kv('   of which no probed resolution', f'=COUNTIF(Plan!E2:E{n},"no resolution*")')
r += 1
s.cell(r, 1, 'Placeholders (ready files that still need a value)').font = bold; r += 1
kv('   "Source" placeholder', f'=COUNTIFS(Plan!A2:A{n},"ready",Plan!H2:H{n},"Source")', 'set via Fix... > Source')
kv('   "Year" placeholder', f'=COUNTIFS(Plan!A2:A{n},"ready",Plan!G2:G{n},"Year")', 'set via Fix... > Year')
r += 1
s.cell(r, 1, 'Flags on ready files').font = bold; r += 1
for k in ['res_mismatch', 'hdr_claimed_but_sdr', 'hdr_uncertain', 'audio_und', 'audio_partly_und']:
    kv(f'   {k}', f'=COUNTIFS(Plan!A2:A{n},"ready",Plan!D2:D{n},"*{k}*")', FLAG_TEXT[k])
r += 1
s.cell(r, 1, 'Token distribution (ready files)').font = bold; r += 1
for label, col, vals in [('Source', 'H', ['Web', 'Rip', 'Source']), ('Resolution', 'I', ['4K', '1440p', '1080p', '720p', '576p', '480p', 'SD']), ('Colour', 'J', ['HDR', 'SDR']), ('Codec', 'K', ['H264', 'HEVC', 'AV1', 'MPEG4', 'VC1', 'MPEG2'])]:
    s.cell(r, 1, f'   {label}').font = font
    for v in vals:
        s.cell(r, 2, v).font = font; s.cell(r, 3, f'=COUNTIFS(Plan!A2:A{n},"ready",Plan!{col}2:{col}{n},"{v}")').font = font; r += 1
r += 1
s.cell(r, 1, 'Sheets').font = bold; r += 1
for a, b in [('Plan', 'every movie file: current name, proposed name, flags, tokens, probed facts. Filter on Status.'), ('Blocked', 'only the files the engine refuses to rename, with the reason.'), ('Flags', 'what each flag means and what to do about it.')]:
    s.cell(r, 1, f'   {a}').font = font; s.cell(r, 2, b).font = grey; r += 1
s.column_dimensions['A'].width = 44; s.column_dimensions['B'].width = 14; s.column_dimensions['C'].width = 90

# ---------------- Blocked sheet ----------------
b = wb.create_sheet('Blocked')
b.append(['Current name', 'Blocked reason', 'Would have been', 'Folder'])
for p in plan:
    if not p.get('ok') and not p.get('unchanged'):
        t = p.get('tokens') or {}
        would = f"{t.get('title','')} ({t.get('year','')}) - {t.get('source','')} {t.get('resolution') or '?'} {t.get('hdr') or '?'} {t.get('codec') or '?'}" if t.get('title') else ''
        b.append([p['from'], p.get('blocked') or '', would, p.get('dir') or ''])
for c in b[1]: c.font = hdr_font; c.fill = hdr_fill
for row in b.iter_rows(min_row=2):
    for c in row: c.font = font
for i, w in enumerate([60, 40, 62, 30], 1): b.column_dimensions[get_column_letter(i)].width = w
b.freeze_panes = 'A2'; b.auto_filter.ref = f'A1:D{b.max_row}'

# ---------------- Flags sheet ----------------
f = wb.create_sheet('Flags')
f.append(['Flag', 'Meaning / what to do'])
for k, v in FLAG_TEXT.items(): f.append([k, v])
for c in f[1]: c.font = hdr_font; c.fill = hdr_fill
for row in f.iter_rows(min_row=2):
    for c in row: c.font = font; c.alignment = Alignment(wrap_text=True, vertical='top')
f.column_dimensions['A'].width = 22; f.column_dimensions['B'].width = 110

wb.save(out)
print('wrote', out, 'rows', n - 1)
